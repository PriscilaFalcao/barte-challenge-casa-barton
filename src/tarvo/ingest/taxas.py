"""Ingestão da planilha de taxas contratadas (02), mantida à mão."""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

from ..core.excecoes import Excecao, TipoExcecao
from ..core.ids import norm_cnpj
from ..core.money import parse_br
from .erros import LayoutError

ARQUIVO = "02_tabela_taxas_contratadas.xlsx"

PRODUTOS_MDR = ("DEBITO", "CREDITO_1X", "CREDITO_2_6X", "CREDITO_7_12X", "PIX")
PRODUTOS = PRODUTOS_MDR + ("ANTECIPACAO",)

_DATA = r"(\d{2}/\d{2}/\d{4})"
_VIGENCIA = [
    (re.compile(rf"^ate {_DATA}$"), "fim"),
    (re.compile(rf"^a partir de {_DATA}$"), "ini"),
    (re.compile(rf"^{_DATA}$"), "ini"),
]


@dataclass(frozen=True)
class ContratoLinha:
    ref_origem: str
    lojista: str
    cnpj14: str
    ec_orbe: str | None
    taxas: dict[str, Decimal | None]  # em pontos percentuais (1.49 = 1,49%)
    spread_remessa: Decimal | None  # None = em branco ou "padrão"
    vigencia_texto: str | None
    vigencia_ini: date | None
    vigencia_fim: date | None
    vigencia_valida: bool
    obs: str | None
    inativo: bool = False


def _norm(texto) -> str:
    s = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def parse_vigencia(texto: str) -> tuple[date | None, date | None]:
    """Gramática: 'até dd/mm/aaaa', 'a partir de dd/mm/aaaa' ou 'dd/mm/aaaa' (início)."""
    t = _norm(texto)
    for padrao, lado in _VIGENCIA:
        m = padrao.match(t)
        if m:
            d = datetime.strptime(m.group(1), "%d/%m/%Y").date()
            return (None, d) if lado == "fim" else (d, None)
    raise ValueError(f"vigência fora da gramática: {texto!r}")


def _pct(celula, onde: str) -> Decimal | None:
    if celula is None or (isinstance(celula, str) and not celula.strip()):
        return None
    if isinstance(celula, str):
        if not celula.strip().endswith("%"):
            raise LayoutError(f"Taxas: {onde}: taxa em texto sem '%': {celula!r}")
        return parse_br(celula).normalize()
    return (Decimal(str(celula)) * 100).normalize()


def _coluna_de(cabecalho: str) -> str | None:
    h = _norm(cabecalho)
    if h == "lojista":
        return "lojista"
    if h == "cnpj":
        return "cnpj"
    if h == "ec orbe":
        return "ec_orbe"
    if h == "debito":
        return "DEBITO"
    if h == "credito a vista":
        return "CREDITO_1X"
    if h.startswith("cred") and "2-6" in h:
        return "CREDITO_2_6X"
    if h.startswith("cred") and "7-12" in h:
        return "CREDITO_7_12X"
    if h == "pix":
        return "PIX"
    if h == "antecipacao":
        return "ANTECIPACAO"
    if h == "spread remessa":
        return "spread"
    if h == "vigencia":
        return "vigencia"
    if h == "obs":
        return "obs"
    return None


OBRIGATORIAS = {
    "lojista": "Lojista", "cnpj": "CNPJ", "ec_orbe": "EC Orbe", "DEBITO": "Débito",
    "CREDITO_1X": "Crédito à vista", "CREDITO_2_6X": "Créd. 2-6x", "CREDITO_7_12X": "Créd. 7-12x",
    "PIX": "Pix", "ANTECIPACAO": "Antecipação", "spread": "Spread remessa", "vigencia": "Vigência",
}


def _aba_ignorada(nome: str) -> bool:
    n = _norm(nome)
    return "antigo" in n or "nao usar" in n


def _achar_cabecalho(wb) -> tuple[object, int, dict[str, int]]:
    for ws in wb.worksheets:
        if _aba_ignorada(ws.title):
            continue
        for i, linha in enumerate(ws.iter_rows(values_only=True), start=1):
            cols = {_coluna_de(c): j for j, c in enumerate(linha) if c is not None and _coluna_de(c)}
            if {"lojista", "cnpj", "ec_orbe"} <= cols.keys():
                faltam = [rot for k, rot in OBRIGATORIAS.items() if k not in cols]
                if faltam:
                    raise LayoutError(f"Taxas: colunas ausentes na aba {ws.title!r}: {', '.join(faltam)}")
                return ws, i, cols
    raise LayoutError("Taxas: cabeçalho (Lojista, CNPJ, EC Orbe) não encontrado em nenhuma aba válida")


def _texto(celula) -> str | None:
    if celula is None:
        return None
    s = str(celula).strip()
    return s or None


def ler_taxas(path: Path) -> tuple[list[ContratoLinha], list[Excecao]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, linha_cab, cols = _achar_cabecalho(wb)
    contratos: list[ContratoLinha] = []
    excecoes: list[Excecao] = []
    for n, linha in enumerate(ws.iter_rows(min_row=linha_cab + 1, values_only=True), start=linha_cab + 1):
        nome = _texto(linha[cols["lojista"]])
        if nome is None or nome.startswith("*"):
            continue
        ref = f"Taxas!L{n}"
        if linha[cols["cnpj"]] is None:
            raise LayoutError(f"Taxas: {ref}: lojista {nome!r} sem CNPJ")
        taxas = {p: _pct(linha[cols[p]], ref) for p in PRODUTOS}
        spread_cru = linha[cols["spread"]]
        spread = None if _norm(spread_cru) in ("", "padrao") else _pct(spread_cru, ref)
        ec = _texto(linha[cols["ec_orbe"]])
        obs = _texto(linha[cols["obs"]]) if "obs" in cols else None
        inativo = obs is not None and "inativo" in _norm(obs)

        vig = linha[cols["vigencia"]]
        ini = fim = None
        valida = False
        vig_texto = None
        if isinstance(vig, (datetime, date)):
            ini, valida = (vig.date() if isinstance(vig, datetime) else vig), True
            vig_texto = ini.strftime("%d/%m/%Y")
        elif _texto(vig) is not None:
            vig_texto = _texto(vig)
            try:
                ini, fim = parse_vigencia(vig_texto)
                valida = True
            except ValueError:
                excecoes.append(Excecao.nova(
                    TipoExcecao.CONTRATO_AMBIGUO, ref, Decimal("0"),
                    f"{nome}: vigência em texto livre que o parser não entende ({vig_texto!r}); linha não usada"))
        elif not inativo:
            excecoes.append(Excecao.nova(
                TipoExcecao.CONTRATO_AMBIGUO, ref, Decimal("0"),
                f"{nome}: linha sem vigência e sem marca de INATIVO; linha não usada"))

        contratos.append(ContratoLinha(
            ref_origem=ref, lojista=nome, cnpj14=norm_cnpj(linha[cols["cnpj"]]),
            ec_orbe=ec, taxas=taxas, spread_remessa=spread, vigencia_texto=vig_texto,
            vigencia_ini=ini, vigencia_fim=fim, vigencia_valida=valida, obs=obs, inativo=inativo,
        ))
    return contratos, excecoes

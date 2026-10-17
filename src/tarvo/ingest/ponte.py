"""Ingestão do e-mail da Ponte Câmbio (03): anexo XLSX com uma linha por operação de remessa."""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from pathlib import Path

import openpyxl

from ..core.ids import norm_cnpj
from ..core.ledger import Ledger
from .erros import LayoutError
from .texto import norm_texto

ARQUIVO = "03_email_ponte_cambio_relatorio_set2026.eml"
ABA = "Operações"

_COLUNAS = {
    "no operacao": "n_operacao", "data contratacao": "data_contratacao", "data liquidacao": "data_liquidacao",
    "cnpj cliente": "cnpj", "cliente": "cliente", "moeda": "moeda", "valor me": "valor_me",
    "taxa cliente": "taxa_cliente", "ptax ref.": "ptax", "valor r$": "valor_brl", "spread r$": "spread_brl",
    "spread %": "spread_pct", "spread estimado r$ (d-1)": "spread_estimado", "iof r$": "iof",
    "status compliance": "status", "observacao": "observacao",
}
_ROTULO = {v: k for k, v in _COLUNAS.items()}
OBRIGATORIAS = ["n_operacao", "data_contratacao", "cnpj", "valor_me", "ptax", "spread_brl", "spread_pct", "status"]
_ROTULO_AMIGAVEL = {"n_operacao": "Nº Operação", "data_contratacao": "Data Contratação", "cnpj": "CNPJ Cliente",
                    "valor_me": "Valor ME", "ptax": "PTAX Ref.", "spread_brl": "Spread R$",
                    "spread_pct": "Spread %", "status": "Status Compliance"}


@dataclass(frozen=True)
class OpPonte:
    n_operacao: str
    data_contratacao: date
    data_liquidacao: date | None
    cnpj14: str
    cliente: str
    moeda: str
    valor_me: Decimal
    ptax: Decimal
    spread_brl: Decimal
    spread_pct: Decimal  # fração, como a Ponte reporta (0.015 = 1,5%)
    status: str
    observacao: str | None
    linha_fisica: int

    @property
    def ref_origem(self) -> str:
        return self.n_operacao


def _dec(v, onde: str) -> Decimal:
    if v is None or isinstance(v, bool):
        raise LayoutError(f"Ponte: {onde}: valor numérico ausente")
    return Decimal(str(v))


def _data(v, onde: str) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return datetime.strptime(str(v).strip(), "%d/%m/%Y").date()
    except ValueError as exc:
        raise LayoutError(f"Ponte: {onde}: data inválida {v!r}") from exc


def _anexo_xlsx(path: Path) -> bytes:
    with open(path, "rb") as f:
        msg = BytesParser(policy=policy.default).parse(f)
    anexos = [p for p in msg.walk() if (p.get_filename() or "").lower().endswith(".xlsx")]
    if len(anexos) != 1:
        raise LayoutError(f"Ponte: esperava exatamente 1 anexo XLSX no e-mail, achei {len(anexos)}")
    return anexos[0].get_payload(decode=True)


def ler_ponte(path: Path, ledger: Ledger | None = None) -> list[OpPonte]:
    wb = openpyxl.load_workbook(io.BytesIO(_anexo_xlsx(path)), data_only=True)
    aba = next((ws for ws in wb.worksheets if norm_texto(ws.title) == norm_texto(ABA)), None)
    if aba is None:
        raise LayoutError(f"Ponte: aba {ABA!r} não encontrada (abas: {', '.join(wb.sheetnames)})")

    linhas = list(aba.iter_rows(values_only=True))
    cab = next((i for i, l in enumerate(linhas) if l and norm_texto(l[0]) == "no operacao"), None)
    if cab is None:
        raise LayoutError("Ponte: cabeçalho 'Nº Operação' não encontrado na aba Operações")
    cols = {_COLUNAS[norm_texto(c)]: j for j, c in enumerate(linhas[cab]) if norm_texto(c) in _COLUNAS}
    faltam = [_ROTULO_AMIGAVEL[c] for c in OBRIGATORIAS if c not in cols]
    if faltam:
        raise LayoutError(f"Ponte: colunas ausentes: {', '.join(faltam)}")

    def cel(linha, campo):
        return linha[cols[campo]] if campo in cols else None

    ops: list[OpPonte] = []
    for i, linha in enumerate(linhas[cab + 1:], start=cab + 2):
        n_op = str(cel(linha, "n_operacao") or "").strip()
        if not n_op.startswith("PC-"):
            # Só o rodapé ("Total de operações") pode ser pulado. Quem tem CNPJ, Valor ME ou Spread R$
            # é operação: se não for reconhecida, falha em vez de sumir da comissão.
            if any(cel(linha, c) not in (None, "") for c in ("cnpj", "valor_me", "spread_brl")):
                raise LayoutError(
                    f"Ponte: linha {i} parece operação (tem CNPJ, Valor ME ou Spread R$) mas o Nº Operação "
                    f"{n_op!r} não começa com PC-")
            if ledger is not None and any(c not in (None, "") for c in linha):
                ledger.ignorar("Ponte", "rodapé ou linha sem dado", f"L{i}")
            continue
        onde = f"linha {i}, {n_op}"
        liq = cel(linha, "data_liquidacao")
        obs = cel(linha, "observacao")
        ops.append(OpPonte(
            n_operacao=n_op,
            data_contratacao=_data(cel(linha, "data_contratacao"), onde),
            data_liquidacao=_data(liq, onde) if liq not in (None, "") else None,
            cnpj14=norm_cnpj(cel(linha, "cnpj")),
            cliente=str(cel(linha, "cliente") or "").strip(),
            moeda=str(cel(linha, "moeda") or "").strip(),
            valor_me=_dec(cel(linha, "valor_me"), onde),
            ptax=_dec(cel(linha, "ptax"), onde),
            spread_brl=_dec(cel(linha, "spread_brl"), onde),
            spread_pct=_dec(cel(linha, "spread_pct"), onde),
            status=str(cel(linha, "status") or "").strip(),
            observacao=str(obs).strip() if obs not in (None, "") else None,
            linha_fisica=i,
        ))
    return ops

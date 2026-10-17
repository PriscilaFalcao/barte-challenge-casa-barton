"""Interface web local para executar o motor e consultar as saídas reais."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import threading
import time
import uuid
import zipfile
from decimal import Decimal
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlsplit

import openpyxl
from .ingest.erros import LayoutError
from .ingest.texto import norm_texto
from .core.ids import norm_cnpj
from .estado.db import Estado, RESOLVIDA
from .ingest.plano_contas import ler_planos
from .output.validators import ValidationError
from .pipeline import ARQUIVOS, _arquivos, fechar

_PROJECT_ROOT = Path(__file__).resolve().parents[2]
_WEB_ROOT = Path(__file__).resolve().parent / "web"
_DATA_ROOT = _PROJECT_ROOT / "data" / "web"
_MAX_UPLOAD_BYTES = 25 * 1024 * 1024
_MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
# Cada campo aceita um ou mais arquivos. O arquivo é guardado com o nome enviado; se faltar o prefixo que o
# pipeline usa para achar a fonte (01_/06_ Orbe, 02_, 03_, 04_, 05a_, 05b_), ele é acrescentado.
_SLOTS = {
    "orbe": (".csv", "01_"),
    "taxas": (".xlsx", "02_"),
    "ponte": (".eml", "03_"),
    "plano": (".csv", "04_"),
    "aurora": (".csv", "05a_"),
    "meridiano": (".ofx", "05b_"),
}


def _stored_name(slot: str, filename: str) -> str:
    prefixes = tuple(p.split("*")[0] for p in ARQUIVOS[slot])
    return filename if filename.startswith(prefixes) else _SLOTS[slot][1] + filename


def _slot_of(file_name: str) -> str | None:
    """Campo a que um arquivo guardado pertence (o inverso de _stored_name), pelo prefixo e pela extensão."""
    suffix = Path(file_name).suffix.lower()
    for slot, (extension, _) in _SLOTS.items():
        if suffix == extension and any(file_name.startswith(p.split("*")[0]) for p in ARQUIVOS[slot]):
            return slot
    return None


def _slot_paths(closure_dir: Path, slot: str) -> list[Path]:
    return _arquivos(closure_dir / "dados", ARQUIVOS[slot])


_TABLES = {
    "comissoes": "comissoes_por_lojista.csv",
    "divergencias": "divergencias_taxa.csv",
    "notas": "notas_a_emitir.csv",
    "lancamentos": "lancamentos_classificados.csv",
    "pendencias": "pendencias.csv",
    "rastro": "rastro.csv",
}
_DOWNLOADS = (*_TABLES.values(), "resumo.md", "execucao.log")
_RUNNING: set[str] = set()
_RUNNING_LOCK = threading.Lock()


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as source:
        return list(csv.DictReader(source, delimiter=";"))


def _cell_value(value: object) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return "" if value is None else str(value)


def _original_row(closure_dir: Path, source_name: str, reference: str) -> dict[str, str] | None:
    source_path = closure_dir / "dados" / source_name  # o rastro guarda o nome do arquivo de origem
    slot = _slot_of(source_name)
    if slot is None or slot == "plano" or not source_path.is_file():
        return None

    if slot == "orbe":
        for row in _read_csv(source_path):
            if row.get("id_transacao") == reference:
                return row
        return None
    if slot == "aurora":
        match = re.fullmatch(r"L(\d+)", reference)
        if match is None:
            return None
        with source_path.open(encoding="utf-8-sig", newline="") as source:
            rows = list(csv.reader(source, delimiter=";"))
        line_number = int(match.group(1))
        if line_number < 1 or line_number > len(rows):
            return None
        header = next(
            (row for row in rows if row and norm_texto(row[0]) == "data"
             and any(norm_texto(value) == "lancamento" for value in row)),
            None,
        )
        if header is None:
            return None
        raw_row = rows[line_number - 1]
        return {"linha_fisica": reference, **{
            name: raw_row[index] if index < len(raw_row) else ""
            for index, name in enumerate(header)
        }}
    if slot == "meridiano":
        text = source_path.read_bytes().decode("cp1252")
        for block in re.split(r"<STMTTRN>", text, flags=re.I)[1:]:
            transaction = re.split(r"</STMTTRN>", block, maxsplit=1, flags=re.I)[0]
            fields = re.findall(r"<([A-Z0-9.]+)>([^<\r\n]*)", transaction, flags=re.I)
            row = {name.upper(): value.strip() for name, value in fields}
            if row.get("FITID") == reference:
                return {"registro_ofx": transaction.strip(), **row}
        return None
    if slot == "taxas":
        match = re.fullmatch(r"Taxas!L(\d+)", reference)
        if match is None:
            return None
        workbook = openpyxl.load_workbook(source_path, data_only=True, read_only=True)
        worksheet = next((ws for ws in workbook.worksheets if norm_texto(ws.title) == "taxas vigentes"), None)
        if worksheet is None:
            worksheet = next((ws for ws in workbook.worksheets if "antigo" not in norm_texto(ws.title)), None)
        if worksheet is None:
            return None
        row_number = int(match.group(1))
        values = next(worksheet.iter_rows(min_row=row_number, max_row=row_number, values_only=True), ())
        return {"planilha": worksheet.title, "linha": str(row_number), **{
            openpyxl.utils.get_column_letter(index): _cell_value(value)
            for index, value in enumerate(values, start=1)
        }}
    message = BytesParser(policy=policy.default).parsebytes(source_path.read_bytes())
    attachments = [part for part in message.walk() if (part.get_filename() or "").lower().endswith(".xlsx")]
    if len(attachments) != 1:
        return None
    workbook = openpyxl.load_workbook(io.BytesIO(attachments[0].get_payload(decode=True)), data_only=True, read_only=True)
    worksheet = next((ws for ws in workbook.worksheets if norm_texto(ws.title) == "operacoes"), None)
    if worksheet is None:
        return None
    rows = worksheet.iter_rows(values_only=True)
    header = None
    header_row = 0
    for index, row in enumerate(rows, start=1):
        if row and norm_texto(row[0]) == "no operacao":
            header = row
            header_row = index
            break
    if header is None:
        return None
    for index, row in enumerate(worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if row and str(row[0]).strip() == reference:
            return {"linha_planilha": str(index), **{
                _cell_value(name): _cell_value(row[column]) if column < len(row) else ""
                for column, name in enumerate(header) if name is not None
            }}
    return None


def _format_brl(value: Decimal) -> str:
    return f"R$ {value:,.2f}".replace(",", "\0").replace(".", ",").replace("\0", ".")


def _summary(output_dir: Path) -> dict:
    commissions = _read_csv(output_dir / _TABLES["comissoes"])
    divergences = _read_csv(output_dir / _TABLES["divergencias"])
    notes = _read_csv(output_dir / _TABLES["notas"])
    bank_entries = _read_csv(output_dir / _TABLES["lancamentos"])
    pending = _read_csv(output_dir / _TABLES["pendencias"])
    trace = _read_csv(output_dir / _TABLES["rastro"])

    totals = {
        line: sum(
            (Decimal(row["comissao_brl"]) for row in commissions if row["linha_negocio"] == line),
            Decimal("0"),
        )
        for line in ("CHECKOUT", "REMESSA")
    }
    trace_totals: dict[str, Decimal] = {}
    for row in trace:
        key = row["id_saida"]
        trace_totals[key] = trace_totals.get(key, Decimal("0")) + Decimal(row["valor"])

    reconciles = all(
        trace_totals.get(row["id_saida"], Decimal("0")) == Decimal(row[value_column])
        for rows, value_column in (
            (commissions, "comissao_brl"),
            (divergences, "diferenca_brl"),
            (notes, "valor_brl"),
            (bank_entries, "valor"),
            (pending, "valor_brl"),
        )
        for row in rows
    )
    by_decider: dict[str, int] = {}
    for row in pending:
        decider = row.get("quem_decide") or "Não definido"
        by_decider[decider] = by_decider.get(decider, 0) + 1

    return {
        "checkout": _format_brl(totals["CHECKOUT"]),
        "remessa": _format_brl(totals["REMESSA"]),
        "divergences_count": len(divergences),
        "divergences_total": _format_brl(
            sum((Decimal(row["diferenca_brl"]) for row in divergences), Decimal("0"))
        ),
        "notes_count": len(notes),
        "notes_total": _format_brl(
            sum((Decimal(row["valor_brl"]) for row in notes), Decimal("0"))
        ),
        "bank_entries_count": len(bank_entries),
        "pending_count": len(pending),
        "pending_by_decider": by_decider,
        "reconciles": reconciles,
    }


def _load_state(closure_id: str) -> tuple[Path, dict]:
    try:
        normalized_id = str(uuid.UUID(closure_id))
    except ValueError as exc:
        raise FileNotFoundError("Fechamento não encontrado.") from exc
    closure_dir = _DATA_ROOT / "fechamentos" / normalized_id
    state_path = closure_dir / "state.json"
    if not state_path.is_file():
        raise FileNotFoundError("Fechamento não encontrado.")
    with state_path.open(encoding="utf-8") as source:
        state = json.load(source)
    state["files"] = {slot: [names] if isinstance(names, str) else names  # estado antigo: um nome por campo
                      for slot, names in state.get("files", {}).items()}
    return closure_dir, state


def _save_state(closure_dir: Path, state: dict) -> None:
    temporary_path = closure_dir / "state.json.tmp"
    with temporary_path.open("w", encoding="utf-8") as target:
        json.dump(state, target, ensure_ascii=False, indent=2)
    temporary_path.replace(closure_dir / "state.json")


def _response_state(closure_dir: Path, state: dict) -> dict:
    response = {
        "id": state["id"],
        "month": state["month"],
        "status": state["status"],
        "files": {slot: names for slot, names in state["files"].items() if names},
        "validation_error": state.get("validation_error"),
        "validated": bool(state.get("validated")),
        "execution_id": state.get("execution_id"),
        "closure_record": state.get("closure_record"),
    }
    result_dir = closure_dir / ("saida" if state["status"] in {"completed", "closed"} else "validacao")
    if state.get("validated") or state["status"] in {"validated", "completed", "closed"}:
        response["summary"] = _summary(result_dir)
        response["outputs"] = list(_TABLES)
    if state.get("error"):
        response["error"] = state["error"]
    response["closed"] = state["status"] == "closed"
    return response


class _Handler(BaseHTTPRequestHandler):
    server_version = "TarvoWeb/1.0"

    def do_GET(self) -> None:
        request = urlsplit(self.path)
        if request.path == "/":
            self._static("index.html", "text/html; charset=utf-8")
            return
        if request.path in {"/app.css", "/app.js"}:
            file_name = request.path.lstrip("/")
            content_type = "text/css; charset=utf-8" if file_name.endswith(".css") else "text/javascript; charset=utf-8"
            self._static(file_name, content_type)
            return
        match = re.fullmatch(r"/api/closures/([0-9a-f-]{36})", request.path)
        if match:
            self._get_closure(match.group(1))
            return
        match = re.fullmatch(r"/api/closures/([0-9a-f-]{36})/(queue|accounts|activity)", request.path)
        if match:
            {"queue": self._get_queue, "accounts": self._get_accounts, "activity": self._get_activity}[
                match.group(2)
            ](match.group(1))
            return
        match = re.fullmatch(r"/api/closures/([0-9a-f-]{36})/events", request.path)
        if match:
            self._stream_events(match.group(1))
            return
        table_names = "|".join(_TABLES)
        match = re.fullmatch(rf"/api/closures/([0-9a-f-]{{36}})/tables/({table_names})", request.path)
        if match:
            self._get_table(match.group(1), match.group(2))
            return
        match = re.fullmatch(r"/api/closures/([0-9a-f-]{36})/trace", request.path)
        if match:
            self._get_trace(match.group(1), parse_qs(request.query))
            return
        match = re.fullmatch(r"/api/closures/([0-9a-f-]{36})/source", request.path)
        if match:
            self._get_source(match.group(1), parse_qs(request.query))
            return
        match = re.fullmatch(r"/api/closures/([0-9a-f-]{36})/download/([^/]+)", request.path)
        if match:
            self._download(match.group(1), match.group(2))
            return
        self._json({"error": "Endereço não encontrado."}, 404)

    def do_DELETE(self) -> None:
        request = urlsplit(self.path)
        slots = "|".join(_SLOTS)
        match = re.fullmatch(rf"/api/closures/([0-9a-f-]{{36}})/files/({slots})", request.path)
        if match is None:
            self._json({"error": "Endereço não encontrado."}, 404)
            return
        self._remove_file(match.group(1), match.group(2), parse_qs(request.query).get("name", [None])[0])

    def _remove_file(self, closure_id: str, slot: str, name: str | None) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        if state["status"] != "uploading":
            self._json({"error": "Os arquivos só podem ser alterados antes da pré-validação."}, 409)
            return
        names = state["files"].get(slot, [])
        if name is not None and name not in names:
            self._json({"error": "Arquivo não encontrado neste campo."}, 404)
            return
        for removed in ([name] if name is not None else list(names)):  # sem nome: esvazia o campo
            (closure_dir / "dados" / _stored_name(slot, removed)).unlink(missing_ok=True)
            names.remove(removed)
        _save_state(closure_dir, state)
        self._json(_response_state(closure_dir, state))

    def do_POST(self) -> None:
        request = urlsplit(self.path)
        if request.path == "/api/closures":
            self._create_closure()
            return
        slots = "|".join(_SLOTS)
        match = re.fullmatch(rf"/api/closures/([0-9a-f-]{{36}})/files/({slots})", request.path)
        if match:
            self._upload_file(match.group(1), match.group(2))
            return
        match = re.fullmatch(r"/api/closures/([0-9a-f-]{36})/validate", request.path)
        if match:
            self._validate(match.group(1))
            return
        match = re.fullmatch(r"/api/closures/([0-9a-f-]{36})/run", request.path)
        if match:
            self._run(match.group(1))
            return
        match = re.fullmatch(r"/api/closures/([0-9a-f-]{36})/decisions", request.path)
        if match:
            self._decide(match.group(1))
            return
        match = re.fullmatch(r"/api/closures/([0-9a-f-]{36})/close", request.path)
        if match:
            self._close_closure(match.group(1))
            return
        self._json({"error": "Endereço não encontrado."}, 404)

    def log_message(self, format: str, *args) -> None:
        print(f"{self.log_date_time_string()} {self.address_string()} {format % args}")

    def _static(self, file_name: str, content_type: str) -> None:
        path = _WEB_ROOT / file_name
        if not path.is_file():
            self._json({"error": "Arquivo da interface não encontrado."}, 500)
            return
        body = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Content-Security-Policy", "default-src 'self'; style-src 'self'; script-src 'self'; connect-src 'self';")
        self.end_headers()
        self.wfile.write(body)

    def _json(self, payload: dict, status: int = 200) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(body)

    def _read_json(self) -> dict:
        length = self._content_length()
        if length > 64 * 1024:
            raise ValueError("A solicitação é maior do que o limite permitido.")
        try:
            payload = json.loads(self.rfile.read(length))
        except json.JSONDecodeError as exc:
            raise ValueError("Envie os dados do formulário em JSON válido.") from exc
        if not isinstance(payload, dict):
            raise ValueError("Os dados enviados precisam ser um objeto.")
        return payload

    def _content_length(self) -> int:
        raw_length = self.headers.get("Content-Length")
        if raw_length is None:
            raise ValueError("A solicitação não informou o tamanho do arquivo.")
        try:
            length = int(raw_length)
        except ValueError as exc:
            raise ValueError("Tamanho de arquivo inválido.") from exc
        if length < 0:
            raise ValueError("Tamanho de arquivo inválido.")
        return length

    def _create_closure(self) -> None:
        try:
            payload = self._read_json()
        except ValueError as exc:
            self._json({"error": str(exc)}, 400)
            return
        month = payload.get("month")
        if not isinstance(month, str) or not _MONTH_RE.fullmatch(month):
            self._json({"error": "Informe o mês no formato AAAA-MM."}, 400)
            return
        closure_id = str(uuid.uuid4())
        closure_dir = _DATA_ROOT / "fechamentos" / closure_id
        (closure_dir / "dados").mkdir(parents=True, exist_ok=False)
        state = {"id": closure_id, "month": month, "status": "uploading", "files": {}}
        _save_state(closure_dir, state)
        self._json(_response_state(closure_dir, state), 201)

    def _upload_file(self, closure_id: str, slot: str) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        if state["status"] != "uploading":
            self._json({"error": "Os arquivos só podem ser alterados antes da pré-validação."}, 409)
            return
        try:
            length = self._content_length()
        except ValueError as exc:
            self._json({"error": str(exc)}, 400)
            return
        if length == 0 or length > _MAX_UPLOAD_BYTES:
            self.close_connection = True
            self._json({"error": "O arquivo precisa ter conteúdo e não pode ultrapassar 25 MB."}, 413)
            return
        filename = unquote(self.headers.get("X-Filename", ""))
        expected_extension = _SLOTS[slot][0]
        if not filename or Path(filename).name != filename or Path(filename).suffix.lower() != expected_extension:
            self._json({"error": f"Este campo aceita somente arquivos {expected_extension}."}, 415)
            return
        body = self.rfile.read(length)
        target = closure_dir / "dados" / _stored_name(slot, filename)
        target.write_bytes(body)
        names = state["files"].setdefault(slot, [])
        if filename not in names:  # o mesmo nome enviado de novo substitui o arquivo
            names.append(filename)
        _save_state(closure_dir, state)
        self._json(_response_state(closure_dir, state))

    def _validate(self, closure_id: str) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        if state["status"] != "uploading":
            self._json({"error": "A pré-validação só pode ser feita durante o envio dos arquivos."}, 409)
            return
        missing = [slot for slot in _SLOTS if not state["files"].get(slot)]
        if missing:
            labels = {
                "orbe": "relatório Orbe",
                "taxas": "planilha de taxas",
                "ponte": "e-mail da Ponte",
                "plano": "plano de contas",
                "aurora": "extrato Aurora",
                "meridiano": "extrato Meridiano",
            }
            self._json({"error": "Envie os arquivos obrigatórios: " + ", ".join(labels[slot] for slot in missing) + "."}, 409)
            return
        try:
            fechar(state["month"], closure_dir / "dados", closure_dir / "validacao", cliente=None)
        except (LayoutError, ValidationError, OSError, ValueError) as exc:
            state["validation_error"] = str(exc)
            state["status"] = "uploading"
            _save_state(closure_dir, state)
            self._json({"error": f"Não foi possível validar os arquivos: {exc}"}, 422)
            return
        state["status"] = "validated"
        state["validated"] = True
        state.pop("validation_error", None)
        _save_state(closure_dir, state)
        self._json(_response_state(closure_dir, state))

    def _run(self, closure_id: str) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        if state["status"] not in {"validated", "completed", "failed"} or not state.get("validated"):
            self._json({"error": "Pré-valide os arquivos antes de executar o fechamento."}, 409)
            return
        if not self._start_execution(closure_dir, state):
            self._json({"error": "Já existe uma execução em andamento para este fechamento."}, 409)
            return
        self._json(_response_state(closure_dir, state), 202)

    def _start_execution(self, closure_dir: Path, state: dict) -> bool:
        closure_id = state["id"]
        if state["status"] == "closed":
            return False
        with _RUNNING_LOCK:
            if closure_id in _RUNNING:
                return False
            _RUNNING.add(closure_id)
        try:
            state["status"] = "running"
            state.pop("error", None)
            state["previous_execution_id"] = state.get("execution_id", 0)
            state.pop("execution_id", None)
            state["started_at"] = datetime.now().isoformat(timespec="seconds")
            _save_state(closure_dir, state)
            worker = threading.Thread(
                target=self._run_worker,
                args=(closure_dir, state),
                name=f"tarvo-web-{closure_id[:8]}",
                daemon=True,
            )
            worker.start()
        except Exception:
            with _RUNNING_LOCK:
                _RUNNING.discard(closure_id)
            raise
        return True

    @staticmethod
    def _run_worker(closure_dir: Path, state: dict) -> None:
        closure_id = state["id"]
        try:
            result = fechar(
                state["month"],
                closure_dir / "dados",
                closure_dir / "saida",
                estado=closure_dir / "estado",
            )
            state["status"] = "completed"
            state["execution_id"] = result.execucao_id
            state.pop("error", None)
        except Exception as exc:
            state["status"] = "failed"
            state["error"] = f"Não foi possível concluir o fechamento: {exc}"
            print(f"Falha na execução web {closure_id}: {type(exc).__name__}: {exc}")
        finally:
            try:
                _save_state(closure_dir, state)
            except OSError as exc:
                print(f"Não foi possível persistir o estado do fechamento {closure_id}: {exc}")
            finally:
                with _RUNNING_LOCK:
                    _RUNNING.discard(closure_id)

    def _get_queue(self, closure_id: str) -> None:
        try:
            _, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        if not (state.get("validated") or state["status"] in {"completed", "closed"}):
            self._json({"error": "A fila fica disponível depois da pré-validação."}, 409)
            return
        db = Estado(_DATA_ROOT / "fechamentos" / state["id"] / "estado")
        try:
            items = db.fila(state["month"])
        finally:
            db.fechar()
        self._json({
            "items": items,
            "open_count": len(items),
            "blocking_count": sum(item["bloqueante"] == "S" for item in items),
            "closed": state["status"] == "closed",
        })

    def _get_accounts(self, closure_id: str) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        if not state["files"].get("plano"):
            self._json({"error": "Envie o plano de contas antes de consultar as contas."}, 409)
            return
        try:
            accounts = ler_planos(_slot_paths(closure_dir, "plano"))
        except (LayoutError, OSError, UnicodeError) as exc:
            self._json({"error": f"Não foi possível ler o plano de contas: {exc}"}, 422)
            return
        self._json({"accounts": [
            {"codigo": account.codigo, "descricao": account.descricao, "analitica": account.analitica}
            for account in accounts.values() if account.analitica
        ]})

    def _get_activity(self, closure_id: str) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        db = Estado(closure_dir / "estado")
        try:
            execution = db.ultima_execucao(state["month"])
            activity = {
                "status": state["status"],
                "error": state.get("error"),
                "execution": execution,
                "events": db.eventos(execution["id"]) if execution else [],
                "agent_calls": db.chamadas(execution["id"]) if execution else [],
                "closed": state["status"] == "closed",
            }
        finally:
            db.fechar()
        self._json(activity)

    def _stream_events(self, closure_id: str) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        if state["status"] not in {"running", "completed", "failed", "closed"}:
            self._json({"error": "Não há execução para acompanhar."}, 409)
            return
        self.send_response(200)
        self.send_header("Content-Type", "text/event-stream; charset=utf-8")
        self.send_header("Cache-Control", "no-cache")
        self.send_header("Connection", "close")
        self.send_header("X-Accel-Buffering", "no")
        self.end_headers()
        self.close_connection = True
        last_event_id = 0
        heartbeat_at = time.monotonic()
        db = Estado(closure_dir / "estado")
        try:
            while True:
                _, current = _load_state(closure_id)
                execution = db.ultima_execucao(current["month"])
                if execution and execution["id"] <= current.get("previous_execution_id", 0):
                    execution = None
                events = db.eventos(execution["id"]) if execution else []
                execution_status = execution["status"] if execution else None
                if execution:
                    for event in events:
                        if event["id"] <= last_event_id:
                            continue
                        payload = json.dumps(event, ensure_ascii=False)
                        self.wfile.write(f"id: {event['id']}\nevent: progress\ndata: {payload}\n\n".encode("utf-8"))
                        self.wfile.flush()
                        last_event_id = event["id"]
                if current["status"] in {"completed", "failed", "closed"} and execution_status != "EXECUTANDO":
                    payload = json.dumps(_response_state(closure_dir, current), ensure_ascii=False)
                    self.wfile.write(f"event: complete\ndata: {payload}\n\n".encode("utf-8"))
                    self.wfile.flush()
                    return
                if time.monotonic() - heartbeat_at >= 10:
                    self.wfile.write(b": keep-alive\n\n")
                    self.wfile.flush()
                    heartbeat_at = time.monotonic()
                time.sleep(0.5)
        except (BrokenPipeError, ConnectionResetError):
            return
        finally:
            db.fechar()

    def _decide(self, closure_id: str) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
            payload = self._read_json()
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        except ValueError as exc:
            self._json({"error": str(exc)}, 400)
            return
        if state["status"] not in {"completed", "failed"} or not state.get("validated"):
            self._json({"error": "As decisões ficam disponíveis após a execução do fechamento."}, 409)
            return
        item_id = payload.get("id_saida")
        decision_type = payload.get("tipo_decisao")
        scope = payload.get("escopo", "caso")
        author = payload.get("autor")
        reason = payload.get("motivo")
        if not all(isinstance(value, str) and value.strip() for value in (item_id, decision_type, scope, author, reason)):
            self._json({"error": "Informe a pendência, a decisão, o autor e o motivo."}, 400)
            return

        db = Estado(closure_dir / "estado")
        try:
            item = db.item_fila(state["month"], item_id)
            if item is None or item["estado"] == RESOLVIDA:
                self._json({"error": "Esta pendência não está mais aberta na fila."}, 404)
                return
            if decision_type == "CONTA":
                account_code = payload.get("conta")
                if not isinstance(account_code, str) or not account_code.strip():
                    self._json({"error": "Selecione uma conta analítica."}, 400)
                    return
                accounts = ler_planos(_slot_paths(closure_dir, "plano"))
                account = accounts.get(account_code)
                if account is None or not account.analitica:
                    self._json({"error": "A conta selecionada não existe ou não aceita lançamentos."}, 422)
                    return
                if scope == "caso" and not item_id.startswith("PEN-CLASSIFICACAO") \
                        and not re.fullmatch(r"(L\d+|MRD\w+)", item["ref_origem"]):
                    self._json({"error": "A decisão de conta por caso só vale para lançamentos bancários."}, 422)
                    return
                if scope not in {"caso", "regra"}:
                    self._json({"error": "O escopo precisa ser caso ou regra."}, 400)
                    return
                decision_payload = {"ref_origem": item["ref_origem"], "conta": account_code}
                if scope == "regra":
                    pattern = payload.get("padrao")
                    bank = payload.get("banco")
                    if not isinstance(pattern, str) or not pattern or bank not in {"AURORA", "MERIDIANO"}:
                        self._json({"error": "Uma regra exige regex e banco Aurora ou Meridiano."}, 400)
                        return
                    try:
                        re.compile(pattern)
                    except re.error as exc:
                        self._json({"error": f"A expressão da regra não é válida: {exc}"}, 400)
                        return
                decision_id = db.registrar_decisao(
                    state["month"], item_id, "CONTA", decision_payload, scope, author, reason
                )
                if scope == "regra":
                    db.promover_regra(
                        "PADRAO_CONTA",
                        {"id": f"HUM-{decision_id}", "banco": bank, "padrao": pattern,
                         "conta": account_code, "status": f"decisão de {author}: {reason}"},
                        author, reason, decision_id,
                    )
            elif decision_type == "EC_LOJISTA":
                ec = payload.get("ec")
                cnpj = payload.get("cnpj")
                if not isinstance(ec, str) or not ec.strip() or not isinstance(cnpj, str) or not cnpj.strip():
                    self._json({"error": "Informe o EC e o CNPJ do lojista."}, 400)
                    return
                try:
                    normalized_cnpj = norm_cnpj(cnpj)
                except ValueError as exc:
                    self._json({"error": str(exc)}, 422)
                    return
                decision_id = db.registrar_decisao(
                    state["month"], item_id, "EC_LOJISTA",
                    {"ec": ec.strip(), "cnpj": normalized_cnpj}, "regra", author, reason,
                )
                db.promover_regra(
                    "EC_LOJISTA", {"ec": ec.strip(), "cnpj": normalized_cnpj}, author, reason, decision_id
                )
            elif decision_type == "NOTA":
                note = payload.get("texto")
                if not isinstance(note, str) or not note.strip():
                    self._json({"error": "Informe o texto da decisão."}, 400)
                    return
                decision_id = db.registrar_decisao(
                    state["month"], item_id, "NOTA", {"texto": note.strip()}, "caso", author, reason
                )
            else:
                self._json({"error": "Tipo de decisão não reconhecido."}, 400)
                return
        except (LayoutError, OSError) as exc:
            self._json({"error": f"Não foi possível validar a decisão: {exc}"}, 422)
            return
        except ValueError as exc:
            self._json({"error": str(exc)}, 400)
            return
        finally:
            db.fechar()

        if not self._start_execution(closure_dir, state):
            self._json({"error": "A decisão foi registrada, mas não foi possível iniciar a reexecução."}, 409)
            return
        self._json({"decision_id": decision_id, "status": "running"}, 202)

    def _close_closure(self, closure_id: str) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
            payload = self._read_json()
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        except ValueError as exc:
            self._json({"error": str(exc)}, 400)
            return
        author = payload.get("autor")
        reason = payload.get("motivo")
        if not isinstance(author, str) or not author.strip() or not isinstance(reason, str) or not reason.strip():
            self._json({"error": "Informe quem está fechando e o motivo."}, 400)
            return
        if state["status"] != "completed":
            self._json({"error": "A execução precisa estar concluída antes de fechar a competência."}, 409)
            return
        pending = _read_csv(closure_dir / "saida" / _TABLES["pendencias"])
        blockers = [row for row in pending if row.get("bloqueante") == "S"]
        if blockers:
            self._json({
                "error": f"Não é possível fechar: ainda há {len(blockers)} pendência(s) bloqueante(s).",
                "blocking_count": len(blockers),
            }, 409)
            return
        hashes = {
            name: hashlib.sha256((closure_dir / "saida" / name).read_bytes()).hexdigest()
            for name in _DOWNLOADS
        }
        state["closure_record"] = {
            "closed_by": author.strip(),
            "reason": reason.strip(),
            "closed_at": datetime.now().isoformat(timespec="seconds"),
            "files_sha256": hashes,
        }
        state["status"] = "closed"
        _save_state(closure_dir, state)
        self._json(_response_state(closure_dir, state))

    def _get_closure(self, closure_id: str) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        self._json(_response_state(closure_dir, state))

    def _get_table(self, closure_id: str, table: str) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        if state["status"] not in {"completed", "closed"}:
            self._json({"error": "As saídas aparecem depois da execução do fechamento."}, 409)
            return
        try:
            rows = _read_csv(closure_dir / "saida" / _TABLES[table])
        except OSError as exc:
            self._json({"error": f"Não foi possível abrir a saída: {exc}"}, 500)
            return
        self._json({"rows": rows})

    def _get_trace(self, closure_id: str, query: dict[str, list[str]]) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        if state["status"] not in {"completed", "closed"}:
            self._json({"error": "O rastro aparece depois da execução do fechamento."}, 409)
            return
        output_id = query.get("id_saida", [""])[0]
        if not output_id or len(output_id) > 128:
            self._json({"error": "Identificador de saída inválido."}, 400)
            return
        try:
            rows = _read_csv(closure_dir / "saida" / _TABLES["rastro"])
        except OSError as exc:
            self._json({"error": f"Não foi possível abrir o rastro: {exc}"}, 500)
            return
        self._json({"rows": [row for row in rows if row["id_saida"] == output_id]})

    def _get_source(self, closure_id: str, query: dict[str, list[str]]) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        if state["status"] not in {"completed", "closed"}:
            self._json({"error": "A linha original fica disponível após a execução."}, 409)
            return
        source_name = query.get("arquivo", [""])[0]
        reference = query.get("ref_origem", [""])[0]
        if not source_name or not reference or len(reference) > 128:
            self._json({"error": "Arquivo de origem ou referência inválidos."}, 400)
            return
        try:
            row = _original_row(closure_dir, source_name, reference)
        except (OSError, ValueError, openpyxl.utils.exceptions.InvalidFileException) as exc:
            self._json({"error": f"Não foi possível abrir a linha original: {exc}"}, 422)
            return
        if row is None:
            self._json({"error": "A linha original não foi encontrada para esta referência."}, 404)
            return
        self._json({"arquivo": source_name, "referencia": reference, "linha": row})

    def _download(self, closure_id: str, name: str) -> None:
        try:
            closure_dir, state = _load_state(closure_id)
        except FileNotFoundError as exc:
            self._json({"error": str(exc)}, 404)
            return
        if state["status"] not in {"completed", "closed"}:
            self._json({"error": "Os arquivos ficam disponíveis depois da execução."}, 409)
            return
        if name == "todos.zip":
            archive_path = closure_dir / "saida" / "todos.zip"
            with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as archive:
                for file_name in _DOWNLOADS:
                    archive.write(closure_dir / "saida" / file_name, arcname=file_name)
            path = archive_path
            content_type = "application/zip"
        elif name in _DOWNLOADS:
            path = closure_dir / "saida" / name
            content_type = "text/markdown; charset=utf-8" if name.endswith(".md") else (
                "text/plain; charset=utf-8" if name.endswith(".log") else "text/csv; charset=utf-8"
            )
        else:
            self._json({"error": "Arquivo de saída não encontrado."}, 404)
            return
        try:
            body = path.read_bytes()
        except OSError as exc:
            self._json({"error": f"Não foi possível abrir a saída: {exc}"}, 500)
            return
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Disposition", f'attachment; filename="{name}"')
        self.send_header("Content-Length", str(len(body)))
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(body)


def serve(port: int = 8000) -> None:
    if not 1 <= port <= 65535:
        raise ValueError("A porta precisa estar entre 1 e 65535.")
    _DATA_ROOT.mkdir(parents=True, exist_ok=True)
    server = ThreadingHTTPServer(("127.0.0.1", port), _Handler)
    print(f"Tarvo disponível em http://127.0.0.1:{port} (Ctrl+C para encerrar)")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nInterface encerrada.")
    finally:
        server.server_close()

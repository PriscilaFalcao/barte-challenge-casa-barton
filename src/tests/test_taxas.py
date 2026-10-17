from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from tarvo.core.excecoes import TipoExcecao
from tarvo.ingest.erros import LayoutError
from tarvo.ingest.taxas import ler_taxas, parse_vigencia

CAB = ["Lojista", "CNPJ", "EC Orbe", "Débito", "Crédito à vista", "Créd. 2-6x", "Créd. 7-12x",
       "Pix", "Antecipação", "Spread remessa", "Vigência", "Obs"]


def _xlsx(tmp_path, linhas, topo=4, cab=CAB, extra_sheets=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Taxas vigentes"
    for i in range(topo - 1):
        ws.append([f"titulo {i}"])
    ws.append([None, None, None, "CARTÃO / PIX (MDR)"])
    ws.append(cab)
    for l in linhas:
        ws.append(l)
    for nome in extra_sheets:
        w2 = wb.create_sheet(nome)
        w2.append(["Lojista", "CNPJ", "Débito"])
        w2.append(["Fantasma", "11111111000111", 0.0169])
    p = tmp_path / "taxas.xlsx"
    wb.save(p)
    return p


L1 = ["Loja A", "22.224.579/0001-09", "1075624862", 0.0149, 0.0259, 0.0319, 0.0379, 0.0089, 0.021, None, "01/01/2026", None]


def test_header_found_by_text_regardless_of_row(tmp_path):
    a, _ = ler_taxas(_xlsx(tmp_path, [L1], topo=4))
    b, _ = ler_taxas(_xlsx(tmp_path, [L1], topo=9))
    assert a[0].cnpj14 == b[0].cnpj14 == "22224579000109"
    assert a[0].ref_origem == "Taxas!L6" and b[0].ref_origem == "Taxas!L11"


def test_rates_become_percent_decimals(tmp_path):
    (c,), _ = ler_taxas(_xlsx(tmp_path, [L1]))
    assert c.taxas["DEBITO"] == Decimal("1.49")
    assert c.taxas["CREDITO_1X"] == Decimal("2.59")
    assert c.taxas["CREDITO_2_6X"] == Decimal("3.19")
    assert c.taxas["CREDITO_7_12X"] == Decimal("3.79")
    assert c.taxas["PIX"] == Decimal("0.89")
    assert c.taxas["ANTECIPACAO"] == Decimal("2.10")
    assert c.ec_orbe == "1075624862"


def test_text_rates_and_numeric_cnpj_without_leading_zero(tmp_path):
    l = ["Loja B", 3229837000103, 1066100982, "1,49%", "2,59%", "3,09%", "3,79%", "0,89%", "2,10%", None, "01/01/2026", None]
    (c,), _ = ler_taxas(_xlsx(tmp_path, [l]))
    assert c.cnpj14 == "03229837000103"
    assert c.ec_orbe == "1066100982"
    assert c.taxas["CREDITO_2_6X"] == Decimal("3.09")


def test_blank_and_padrao_spread_are_none_and_number_is_percent(tmp_path):
    l_padrao = [*L1[:9], "padrão", "01/01/2026", None]
    l_num = [*L1[:9], 0.012, "01/01/2026", None]
    cs, _ = ler_taxas(_xlsx(tmp_path, [L1, l_padrao, l_num]))
    assert [c.spread_remessa for c in cs] == [None, None, Decimal("1.2")]


def test_merchant_without_mdr_has_none_rates(tmp_path):
    l = ["So remessa", "95739378000118", None, None, None, None, None, None, None, None, "01/01/2026", "só remessa"]
    (c,), _ = ler_taxas(_xlsx(tmp_path, [l]))
    assert c.ec_orbe is None
    assert all(v is None for v in c.taxas.values())


def test_footer_note_row_is_not_a_merchant(tmp_path):
    cs, _ = ler_taxas(_xlsx(tmp_path, [L1, ["* spread em branco = padrão"]]))
    assert len(cs) == 1


def test_old_sheet_is_ignored(tmp_path):
    cs, _ = ler_taxas(_xlsx(tmp_path, [L1], extra_sheets=["ANTIGO 2025 - nao usar"]))
    assert [c.lojista for c in cs] == ["Loja A"]


@pytest.mark.parametrize(
    "texto,ini,fim",
    [
        ("até 14/09/2026", None, date(2026, 9, 14)),
        ("a partir de 15/09/2026", date(2026, 9, 15), None),
        ("01/01/2026", date(2026, 1, 1), None),
        ("  A PARTIR DE 15/09/2026 ", date(2026, 9, 15), None),
    ],
)
def test_parse_vigencia(texto, ini, fim):
    assert parse_vigencia(texto) == (ini, fim)


def test_parse_vigencia_rejects_free_text():
    with pytest.raises(ValueError):
        parse_vigencia("depois do carnaval")


def test_unparseable_vigencia_raises_exception_and_line_is_unusable(tmp_path):
    l = [*L1[:10], "depois do carnaval", None]
    (c,), exc = ler_taxas(_xlsx(tmp_path, [l]))
    assert not c.vigencia_valida
    assert [e.tipo for e in exc] == [TipoExcecao.CONTRATO_AMBIGUO]
    assert exc[0].ref_origem == "Taxas!L6" and exc[0].quem_decide == "Comercial"


def test_inactive_without_vigencia_is_unusable_without_exception(tmp_path):
    l = [*L1[:10], None, "INATIVO desde jul/26"]
    (c,), exc = ler_taxas(_xlsx(tmp_path, [l]))
    assert c.inativo and not c.vigencia_valida
    assert exc == []


def test_active_without_vigencia_raises_exception(tmp_path):
    l = [*L1[:10], None, None]
    _, exc = ler_taxas(_xlsx(tmp_path, [l]))
    assert [e.tipo for e in exc] == [TipoExcecao.CONTRATO_AMBIGUO]


def test_missing_rate_column_fails_naming_it(tmp_path):
    cab = [c for c in CAB if c != "Pix"]
    l = [v for v, c in zip(L1, CAB) if c != "Pix"]
    with pytest.raises(LayoutError, match="Pix"):
        ler_taxas(_xlsx(tmp_path, [l], cab=cab))


def test_no_header_fails(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["nada", "aqui"])
    p = tmp_path / "x.xlsx"
    wb.save(p)
    with pytest.raises(LayoutError):
        ler_taxas(p)


def test_real_file_golden(dados):
    cs, exc = ler_taxas(dados / "02_tabela_taxas_contratadas.xlsx")
    assert len(cs) == 25  # linhas 5 a 29 da planilha
    assert exc == []
    bridge = [c for c in cs if c.cnpj14 == "22224579000109"]
    assert [(c.vigencia_ini, c.vigencia_fim) for c in bridge] == [
        (None, date(2026, 9, 14)), (date(2026, 9, 15), None)]
    assert bridge[0].taxas["CREDITO_1X"] == Decimal("2.59") and bridge[1].taxas["CREDITO_1X"] == Decimal("2.19")
    assert {c.lojista for c in cs if c.inativo} == {"Globo Trotter Viagens Ltda"}


def test_row_with_data_but_no_merchant_name_fails_loudly(tmp_path):
    sem_nome = [None, *L1[1:]]
    with pytest.raises(LayoutError, match="Taxas!L7.*nome"):
        ler_taxas(_xlsx(tmp_path, [L1, sem_nome]))


def test_row_with_only_cnpj_and_no_name_fails_too(tmp_path):
    with pytest.raises(LayoutError, match="Taxas!L6"):
        ler_taxas(_xlsx(tmp_path, [[None, "22.224.579/0001-09"]]))


def test_footer_is_counted_in_the_ledger_and_blank_rows_are_not(tmp_path):
    from tarvo.core.ledger import Ledger

    ledger = Ledger()
    ler_taxas(_xlsx(tmp_path, [L1, [None], ["* spread em branco = padrão"]]), ledger)
    assert ledger.contar("Taxas", "rodapé ou linha sem dado") == 1


def test_real_file_ledger_has_only_the_footer_note(dados):
    from tarvo.core.ledger import Ledger

    ledger = Ledger()
    ler_taxas(dados / "02_tabela_taxas_contratadas.xlsx", ledger)
    assert ledger.contar("Taxas", "rodapé ou linha sem dado") == 1

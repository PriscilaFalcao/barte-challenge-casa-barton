"""Recuperação de meses anteriores pelo histórico de fechamentos (✅ Helena 11h41, caso Atlas)."""
import csv
import shutil
from decimal import Decimal as D

import openpyxl

from tarvo.estado.historico import linhas_para_recuperar
from tarvo.pipeline import ARQUIVOS, fechar

ATLAS_EC, ATLAS_CNPJ = "1076801732", "39566044000119"


def _ler(p):
    with open(p, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def _outubro(dados, destino):
    """Orbe de 'outubro' = a de setembro com datas e ids trocados; a planilha ganha o contrato da Atlas."""
    destino.mkdir(parents=True)
    texto = (dados / "01_orbe_relatorio_comissoes_set2026.csv").read_text(encoding="utf-8")
    texto = texto.replace("/09/2026", "/10/2026").replace("TX88", "TY88").replace("AN88", "AM88")
    (destino / "01_orbe_relatorio_comissoes_out2026.csv").write_text(texto, encoding="utf-8", newline="")
    for nome in ("03_email_ponte_cambio_relatorio_set2026.eml", "04_plano_de_contas_fenix_erp.csv",
                 "05a_extrato_banco_aurora_cc4471_set2026.csv", "05b_extrato_banco_meridiano_cc90233_set2026.ofx"):
        shutil.copy(dados / nome, destino / nome)
    wb = openpyxl.load_workbook(dados / "02_tabela_taxas_contratadas.xlsx")
    ws = wb["Taxas vigentes"]
    ultima = max(r for r in range(6, ws.max_row + 1) if ws.cell(r, 1).value and ws.cell(r, 3).value)
    ws.insert_rows(ultima + 1)
    linha = ["Atlas Idiomas Ltda", ATLAS_CNPJ, ATLAS_EC, 0.0149, 0.0269, 0.0319, 0.0379, 0.0089, 0.021, None,
             "10/09/2026", "contrato confirmado pelo comercial"]
    for c, v in enumerate(linha, start=1):
        ws.cell(ultima + 1, c, v)
    wb.save(destino / "02_tabela_taxas_contratadas.xlsx")


def test_confirmed_contract_bills_the_pending_month_once(dados, tmp_path):
    hist = tmp_path / "hist"
    (hist / "set").mkdir(parents=True)
    shutil.copytree(dados, hist / "set" / "dados")
    fechar("2026-09", dados, hist / "set" / "saida")
    assert any(p["tipo"] == "SEM_CONTRATO" for p in _ler(hist / "set" / "saida" / "pendencias.csv"))

    txs, meses = linhas_para_recuperar(hist, "2026-10", ARQUIVOS["orbe"])
    assert {t.cod_estabelecimento for t in txs} == {ATLAS_EC} and len(txs) == 57
    assert meses == {ATLAS_EC: frozenset({"2026-09"})}
    assert linhas_para_recuperar(hist, "2026-09", ARQUIVOS["orbe"]) == ([], {})  # só meses anteriores

    _outubro(dados, hist / "out" / "dados")
    fechar("2026-10", hist / "out" / "dados", hist / "out" / "saida", historico=hist)
    saida = hist / "out" / "saida"
    rastro = _ler(saida / "rastro.csv")
    recuperadas = [r for r in rastro if r["regra_id"].endswith("+RECUP") and r["id_saida"].startswith("COM-")]
    assert len(recuperadas) == 57 and {r["id_saida"] for r in recuperadas} == {f"COM-CHK-{ATLAS_CNPJ}-202610"}
    nota = next(n for n in _ler(saida / "notas_a_emitir.csv") if n["tomador_cnpj"] == ATLAS_CNPJ)  # entra na nota
    atlas = next(c for c in _ler(saida / "comissoes_por_lojista.csv") if c["cnpj"] == ATLAS_CNPJ)
    soma = sum((D(r["valor"]) for r in rastro if r["id_saida"] == atlas["id_saida"]), D("0"))
    assert soma == D(atlas["comissao_brl"]) > 0
    assert sum(D(r["valor"]) for r in recuperadas) * 2 == D(atlas["comissao_brl"])  # setembro + outubro, mesma base
    assert "Recuperação: 57 linha(s) da Orbe de 2026-09" in (saida / "execucao.log").read_text(encoding="utf-8")
    assert not any(p["tipo"] == "SEM_CONTRATO" for p in _ler(saida / "pendencias.csv"))
    # a Atlas já vendeu em setembro: em outubro não é mais "lojista nova", então não há alerta de antecipação
    assert not any(p["tipo"] == "ALERTA_ANTECIPACAO" for p in _ler(saida / "pendencias.csv"))

    # já faturado: o mês seguinte não puxa de novo
    assert linhas_para_recuperar(hist, "2026-11", ARQUIVOS["orbe"]) == ([], {})


def test_still_without_contract_nothing_is_recovered(dados, tmp_path):
    hist = tmp_path / "hist"
    (hist / "set").mkdir(parents=True)
    shutil.copytree(dados, hist / "set" / "dados")
    fechar("2026-09", dados, hist / "set" / "saida")
    _outubro(dados, hist / "out" / "dados")
    shutil.copy(dados / "02_tabela_taxas_contratadas.xlsx", hist / "out" / "dados" / "02_tabela_taxas_contratadas.xlsx")
    fechar("2026-10", hist / "out" / "dados", hist / "out" / "saida", historico=hist)
    rastro = _ler(hist / "out" / "saida" / "rastro.csv")
    assert not any(r["regra_id"].endswith("+RECUP") for r in rastro)
    pend = [p for p in _ler(hist / "out" / "saida" / "pendencias.csv") if p["tipo"] == "SEM_CONTRATO"]
    assert len(pend) == 1  # a pendência de outubro, só com as linhas de outubro
